"""Place uploaded score tables and graphs into an exported Word report.

The branded templates are real prior reports: each test (WISC-V, WRAML-3, ...)
has its own block made of a heading, boilerplate, *blank* score tables, and
sometimes a literal "(insert WISC-V graph here)" placeholder. For every test the
user attached results for, we find that test's block, swap its blank score
tables / graph placeholder for the uploaded ones, and leave the rest of the
block (descriptions, notes) alone. A test the template has no block for gets a
labelled block of its own instead, so nothing the user attached is ever dropped.

Input is a list of ``{"test_name": str, "mode": "fill"|"insert", "items": [...]}``
where each item is ``{"kind": "table", "caption", "rows", "merges"?, "widths"?}`` or
``{"kind": "image", "caption", "content_type", "data_b64"}`` (the shape produced
by ``score_report_extractor``).

Two ways to place a test's tables:

* ``fill`` (the default): the template's own score tables stay exactly as they are
  and only their blank score cells are filled in (``score_fill``); the template is
  what the clinicians use, so its look is never changed. If nothing in the
  template's tables could be matched, the uploaded tables are inserted instead.
* ``insert``: the template's blank tables are replaced by the uploaded tables as
  they appear in the score report.

Graphs always replace the template's "(insert ... graph here)" placeholder.
"""
from __future__ import annotations

import base64
import re
from collections import Counter
from io import BytesIO
from typing import Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt

from services.score_fill import cell_text, fill_template_tables, header_field, parse_source_tables, summarize

P, TBL, SECT_PR = qn("w:p"), qn("w:tbl"), qn("w:sectPr")
HEADING_MAX_CHARS = 100
HEADER_FILL = "D9E2F3"

# Test families used to tell where one test's block ends and the next begins.
KNOWN_FAMILIES = frozenset(
    {
        "wppsi", "wisc", "wais", "wasi", "wiat", "ktea", "wraml", "cvlt", "ctopp", "nepsy", "dkefs",
        "brief", "basc", "abas", "vineland", "conners", "beery", "vmi", "gort", "towl", "woodcock",
        "ppvt", "evt", "bracken", "ados", "scq", "asrs", "masc", "cdi", "scared", "pai", "leiter",
        "mullen", "bayley", "das", "tops", "sldt", "reel", "sib", "wsr", "aseba", "cbcl", "dabs",
        "rey", "tvcf", "srs", "wj",
    }
)
# Where a block with no home in the template is inserted (before the first of
# these headings that comes after the results begin; otherwise at the end).
RESULTS_START_HEADINGS = ("tests administered", "test results", "general cognitive abilities")
RESULTS_END_PREFIXES = (
    "summary of assessment results", "summary & conclusions", "summary and conclusions", "appendix 1",
)
RESULTS_END_EXACT = ("highlighted recommendations", "recommendations", "additional recommendations")

_PLACEHOLDER = re.compile(r"^\(?\s*insert\b.*\b(graph|chart|figure)\b.*\)?$", re.IGNORECASE)
_ROMAN = {"i": "1", "ii": "2", "iii": "3", "iv": "4", "v": "5", "vi": "6"}


# --------------------------------------------------------------------------- #
# public API
# --------------------------------------------------------------------------- #
def apply_test_results(
    docx_stream: BytesIO,
    test_results: Optional[Sequence[Mapping]],
    section_targets: Sequence[Tuple[str, Sequence[str]]] = (),
    report: Optional[List[dict]] = None,
) -> BytesIO:
    """``report``, if given, receives one dict per test saying what was done
    (``result``: "filled" / "replaced" / "added", plus fill details)."""
    groups = _merge_groups(test_results or [])
    if not groups:
        docx_stream.seek(0)
        return docx_stream

    document = Document(docx_stream)
    body = document.element.body
    font = _body_font(body)
    boundary_headings = {_norm(alias) for _, aliases in section_targets for alias in aliases}
    for name, items, mode in groups:
        outcome = _place_in_template_block(document, body, name, items, boundary_headings, font, mode)
        if not outcome:
            _append_block(document, body, name, items, font)
            outcome = {"result": "added"}
        if report is not None:
            report.append({"test_name": name, **outcome})

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


# --------------------------------------------------------------------------- #
# matching a test name to text in the template
# --------------------------------------------------------------------------- #
def _norm(text: str) -> str:
    """Lowercase, unify dashes, turn trailing roman numerals into digits, drop
    hyphens ("WISC-V" -> "wisc5", "D-KEFS" -> "dkefs")."""
    text = (text or "").lower().replace("‐", "-").replace("‑", "-")
    text = text.replace("‒", "-").replace("–", "-").replace("—", "-").replace("−", "-")
    text = re.sub(
        r"([a-z]{3,})[-\s](iii|ii|iv|vi|v|i)(?![a-z0-9])",
        lambda m: m.group(1) + _ROMAN[m.group(2)],
        text,
    )
    return re.sub(r"\s+", " ", text.replace("-", "")).strip()


def _family(test_name: str) -> str:
    first = re.split(r"[\s/:(),]+", (test_name or "").strip())[0]
    parts = first.split("-")
    family = parts[0] if len(parts[0]) >= 3 else "".join(parts[:2])
    return re.sub(r"[^a-z0-9]", "", family.lower())


def _has_family(norm_text: str, family: str) -> bool:
    return bool(family) and re.search(r"(?<![a-z0-9])" + re.escape(family), norm_text) is not None


def _has_other_test(raw_text: str, family: str) -> bool:
    """Does this text name a different known test? Checked against the text with
    hyphens removed ("D-KEFS" -> "dkefs") and as written ("CVLT-C": the hyphen
    keeps "cvlt" from fusing with the "c" and failing the word-end check)."""
    hyphenless = _norm(raw_text)
    as_written = re.sub(r"\s+", " ", (raw_text or "").lower())
    return any(
        other != family
        and (
            re.search(r"(?<![a-z0-9])" + re.escape(other) + r"(?![a-z])", hyphenless)
            or re.search(r"(?<![a-z0-9])" + re.escape(other) + r"(?![a-z])", as_written)
        )
        for other in KNOWN_FAMILIES
    )


def _text(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()


def _is_bold(paragraph) -> bool:
    for run in paragraph.iter(qn("w:r")):
        properties = run.find(qn("w:rPr"))
        if properties is not None and properties.find(qn("w:b")) is not None and _text(run):
            return True
    return False


def _body_font(body) -> Optional[str]:
    """The font the template's own text is set in. Its document default is often
    Times New Roman while the prose carries an explicit font (e.g. Calibri), so
    inserted text must copy the explicit one or it looks out of place."""
    counts = Counter(
        fonts.get(qn("w:ascii"))
        for fonts in body.iter(qn("w:rFonts"))
        if fonts.get(qn("w:ascii"))
    )
    return counts.most_common(1)[0][0] if counts else None


def _set_font(run, font: Optional[str]) -> None:
    if font:
        run.font.name = font


def _merge_groups(test_results: Iterable[Mapping]) -> List[Tuple[str, List[Mapping], str]]:
    merged: Dict[str, Tuple[str, List[Mapping], str]] = {}
    for group in test_results:
        name = (group.get("test_name") or "").strip()
        items = list(group.get("items") or [])
        if not name or not items:
            continue
        mode = "insert" if group.get("mode") == "insert" else "fill"
        key = _norm(name)
        if key in merged:
            merged[key][1].extend(items)
        else:
            merged[key] = (name, items, mode)
    return list(merged.values())


# --------------------------------------------------------------------------- #
# placing into the template's own block for the test
# --------------------------------------------------------------------------- #
def _place_in_template_block(
    document, body, name: str, items: List[Mapping], boundary_headings, font: Optional[str] = None,
    mode: str = "fill",
) -> Optional[dict]:
    """Returns what was done ({"result": ...}) or None if the template has no block for the test."""
    children = list(body)
    family = _family(name)
    exact = _norm(name).replace(" ", "")
    has_tables = any(i.get("kind") == "table" for i in items)
    has_images = any(i.get("kind") == "image" for i in items)

    if mode == "fill" and has_tables:
        outcome = _fill_tagged_tables(document, body, name, items, font)
        if outcome:
            return outcome

    candidates: Dict[str, List[int]] = {"eh": [], "fh": [], "ea": [], "fa": []}
    for index, element in enumerate(children):
        if element.tag != P:
            continue
        raw = _text(element)
        norm = _norm(raw)
        if not norm:
            continue
        is_exact = exact in norm.replace(" ", "")
        if not (is_exact or _has_family(norm, family)):
            continue
        head = len(raw) <= HEADING_MAX_CHARS
        candidates[("e" if is_exact else "f") + ("h" if head else "a")].append(index)

    heading_anchors = set(candidates["eh"] + candidates["fh"])  # real headings, not a sentence that mentions the test
    for anchor in candidates["eh"] + candidates["fh"] + candidates["ea"] + candidates["fa"]:
        end = _block_end(children, anchor, family, boundary_headings)
        block = children[anchor:end]
        shells = [el for el in block if el.tag == TBL]
        placeholders = [el for el in block if el.tag == P and _PLACEHOLDER.match(_text(el) or "x")]
        if not ((has_tables and shells) or (has_images and placeholders)):
            continue  # e.g. the "Tests Administered" list: no blank tables to replace

        if mode == "fill" and has_tables and shells and anchor in heading_anchors:
            # never write into a table that is headed as a different test version
            token = _norm(name).replace(" ", "")
            other_versions = {id(t) for t, tag in _table_groups(body) if tag not in (None, token)}
            usable = [t for t in shells if id(t) not in other_versions]
            outcome = _fill_block(document, block, usable, placeholders, items, font) if usable else None
            if outcome:
                return outcome

        remove: List = []
        if has_tables:
            remove += shells + _labels_for(shells, children, anchor)
        if has_images:
            remove += placeholders

        order = {id(el): i for i, el in enumerate(children)}
        remove = list({id(el): el for el in remove}.values())
        first = min(remove, key=lambda el: order[id(el)])
        marker = OxmlElement("w:p")
        first.addprevious(marker)
        for element in remove:
            element.getparent().remove(element)
        _insert_items(document, marker, items, heading=None, font=font)
        marker.getparent().remove(marker)
        return {"result": "replaced"}
    return None


def _fill_block(document, block, shells, placeholders, items: List[Mapping], font: Optional[str]) -> Optional[dict]:
    """Fill the template's own score tables from the uploaded tables, leaving their
    look alone. Graphs (and uploaded tables that aren't score tables, so can't be
    matched to template rows) go where the template says, or after its last table.
    Returns None, having changed nothing, when no template cell could be filled."""
    return _fill_and_place(document, shells, placeholders, items, font)


def _fill_and_place(document, shells, placeholders, items: List[Mapping], font: Optional[str]) -> Optional[dict]:
    table_items = [i for i in items if i.get("kind") == "table"]
    records = parse_source_tables(table_items)
    if not records:
        return None
    stats = fill_template_tables(shells, records)
    if not stats.filled:
        return None

    extras = [
        i for i in items
        if i.get("kind") != "table" or not parse_source_tables([i])
    ]
    if extras:
        has_images = any(i.get("kind") == "image" for i in extras)
        marker = OxmlElement("w:p")
        if placeholders and has_images:
            placeholders[0].addprevious(marker)
            for placeholder in placeholders:
                placeholder.getparent().remove(placeholder)
        else:
            shells[-1].addnext(marker)
        _insert_items(document, marker, extras, heading=None, font=font)
        marker.getparent().remove(marker)
    return {"result": "filled", **summarize(stats)}


_TOKEN = re.compile(r"(?<![a-z0-9])([a-z]{2,8})(\d+)(?![a-z0-9])")


def _test_token(text: str) -> Optional[str]:
    """The test-and-version a piece of text names ("WISC-V Index" -> "wisc5"), if it
    names a known one."""
    for match in _TOKEN.finditer(_norm(text)):
        if match.group(1) in KNOWN_FAMILIES:
            return match.group(1) + match.group(2)
    return None


def _table_groups(body) -> List[Tuple[object, Optional[str]]]:
    """Every score table in the body with the test version it belongs to (or None).
    A table headed "WISC-V Index" starts that version's group; the tables after it
    ("VCI Subtests" ...) belong to it until another version's table, or a new "Index"
    table that names no version, starts the next group. Needs no heading text, which
    a generated report may have replaced."""
    groups: List[Tuple[object, Optional[str]]] = []
    current: Optional[str] = None
    width = 0
    for table in body.iterchildren(TBL):
        rows = table.findall(qn("w:tr"))
        first = rows[0].findall(qn("w:tc")) if rows else []
        header = cell_text(first[0]) if first else ""
        if not any(header_field(cell_text(tc)) for tc in first):
            # No column headings: a supplementary-subtests table continues the group
            # above when it has the same columns; anything else ends the group.
            if current is not None and len(first) == width:
                groups.append((table, current))
            else:
                current = None
                groups.append((table, None))
            continue
        width = len(first)
        tagged = _test_token(header)
        if tagged:
            current = tagged
        elif _norm(header) in ("index", "indexes", "indices") or not header:
            current = None  # a new group that names no version
        groups.append((table, current))
    return groups


def _owned_tables(body, token: str) -> List:
    return [table for table, tag in _table_groups(body) if tag == token]


def _fill_tagged_tables(document, body, name: str, items: List[Mapping], font: Optional[str]) -> Optional[dict]:
    token = _norm(name).replace(" ", "")
    if not re.fullmatch(r"[a-z]+\d+", token):
        return None
    shells = _owned_tables(body, token)
    if not shells:
        return None
    family = _family(name)
    placeholders = [
        el for el in body.iterchildren(P)
        if _PLACEHOLDER.match(_text(el) or "x") and _has_family(_norm(_text(el)), family)
    ]
    return _fill_and_place(document, shells, placeholders, items, font)


def _block_end(children, anchor: int, family: str, boundary_headings) -> int:
    for index in range(anchor + 1, len(children)):
        element = children[index]
        if element.tag != P:
            continue
        raw = _text(element)
        if not raw:
            continue
        norm = _norm(raw)
        if norm in boundary_headings:
            return index
        if len(raw) <= HEADING_MAX_CHARS and _has_other_test(raw, family) and not _has_family(norm, family):
            return index
    return len(children)


def _labels_for(shells, children, anchor: int) -> List:
    """Bold one-line captions sitting directly above a blank table
    ("WISC-V Index Scores Summary") go away with the table they label."""
    position = {id(el): i for i, el in enumerate(children)}
    labels = []
    for table in shells:
        previous = table.getprevious()
        while previous is not None and previous.tag == P and not _text(previous):
            previous = previous.getprevious()
        if (
            previous is not None
            and previous.tag == P
            and position.get(id(previous), -1) >= anchor
            and _is_bold(previous)
            and len(_text(previous)) <= HEADING_MAX_CHARS
        ):
            labels.append(previous)
    return labels


# --------------------------------------------------------------------------- #
# fallback: a labelled block of its own
# --------------------------------------------------------------------------- #
def _append_block(document, body, name: str, items: List[Mapping], font: Optional[str] = None) -> None:
    marker = OxmlElement("w:p")
    insert_before = _fallback_position(body)
    if insert_before is not None:
        insert_before.addprevious(marker)
    else:
        sect_pr = body.find(SECT_PR)
        if sect_pr is not None:
            sect_pr.addprevious(marker)
        else:
            body.append(marker)
    _insert_items(document, marker, items, heading=name, font=font)
    marker.getparent().remove(marker)


def _fallback_position(body):
    """Just before the summary/recommendations that follow the results, so a
    test the template has no block for still lands with the other results."""
    started = False
    for element in body:
        if element.tag != P:
            continue
        norm = _norm(_text(element))
        if not started:
            started = norm in RESULTS_START_HEADINGS
        elif norm.startswith(RESULTS_END_PREFIXES) or norm in RESULTS_END_EXACT:
            return element
    return None


# --------------------------------------------------------------------------- #
# building the Word content
# --------------------------------------------------------------------------- #
def _insert_items(
    document, marker, items: List[Mapping], heading: Optional[str], font: Optional[str] = None
) -> None:
    if heading:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(heading)
        run.bold = True
        run.font.size = Pt(12)
        _set_font(run, font)
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.keep_with_next = True
        marker.addprevious(paragraph._p)

    for item in items:
        caption = (item.get("caption") or "").strip()
        if caption and caption.lower() != (heading or "").lower():
            paragraph = document.add_paragraph()
            run = paragraph.add_run(caption)
            run.bold = True
            run.font.size = Pt(10)
            _set_font(run, font)
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.keep_with_next = True
            marker.addprevious(paragraph._p)

        if item.get("kind") == "table" and item.get("rows"):
            marker.addprevious(_build_table(document, item["rows"], font, item.get("merges"), item.get("widths")))
            spacer = document.add_paragraph()  # Word merges adjacent tables without a paragraph between
            spacer.paragraph_format.space_after = Pt(4)
            marker.addprevious(spacer._p)
        elif item.get("kind") == "image" and item.get("data_b64"):
            picture = _build_picture(document, item)
            if picture is not None:
                marker.addprevious(picture)


def _header_rows(rows: Sequence[Sequence[str]], merges) -> Tuple[set, Optional[int]]:
    """(title rows, header row). Score reports often open with a one-cell title row
    ("Index Scores") above the real column headings; shading that row as the header
    is wrong, so title rows are only bolded and the shaded header is the first row
    that has several cells filled."""
    columns = max(len(r) for r in rows)
    titles = set()
    for r in range(min(len(rows), 3)):
        filled = sum(1 for c in rows[r] if str(c).strip())
        full_span = any(m[0] == r and m[1] == 0 and m[3] >= columns for m in merges)
        if filled >= 2 and not full_span:
            return titles, r
        titles.add(r)
    return set(), 0


def _build_table(
    document, rows: Sequence[Sequence[str]], font: Optional[str] = None,
    merges: Optional[Sequence[Sequence[int]]] = None, widths: Optional[Sequence[int]] = None,
):
    merges = [m for m in (merges or []) if len(m) == 4]
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    properties = table._tbl.tblPr

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "808080")):
            border.set(qn(f"w:{key}"), value)
        borders.append(border)
    # tblPr children must stay in schema order: borders go before layout/margins/look
    later = next(
        (properties.find(qn(tag)) for tag in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
         if properties.find(qn(tag)) is not None),
        None,
    )
    if later is not None:
        later.addprevious(borders)
    else:
        properties.append(borders)

    width = properties.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:type"), "pct")
    width.set(qn("w:w"), "5000")

    covered = set()
    for r, c, rowspan, colspan in merges:
        if r + rowspan > len(rows) or c + colspan > columns:
            continue
        table.cell(r, c).merge(table.cell(r + rowspan - 1, c + colspan - 1))
        covered.update((r + dr, c + dc) for dr in range(rowspan) for dc in range(colspan) if dr or dc)

    if widths and len(widths) == columns and all(w > 0 for w in widths):
        _apply_widths(document, table, widths)

    title_rows, header_row = _header_rows(rows, merges)
    for r, row in enumerate(rows):
        for c in range(columns):
            if (r, c) in covered:
                continue
            cell = table.cell(r, c)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            lines = str(row[c] if c < len(row) else "").split("\n")
            run = paragraph.add_run(lines[0])
            for line in lines[1:]:
                run.add_break()
                run.add_text(line)
            run.font.size = Pt(9)
            _set_font(run, font)
            if r in title_rows:
                run.bold = True
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif r == header_row:
                run.bold = True
                shading = OxmlElement("w:shd")
                shading.set(qn("w:val"), "clear")
                shading.set(qn("w:color"), "auto")
                shading.set(qn("w:fill"), HEADER_FILL)
                cell._tc.get_or_add_tcPr().append(shading)
    if header_row is not None:
        # repeat the column headings (and any title above them) if the table spans pages
        for r in sorted(title_rows | {header_row}):
            table.rows[r]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    return table._tbl


def _apply_widths(document, table, widths: Sequence[int]) -> None:
    """Keep the score report's relative column widths, scaled to the page's text width."""
    try:
        section = document.sections[-1]
        total = int((section.page_width - section.left_margin - section.right_margin) / 635)  # EMU -> twips
    except Exception:
        total = 9360
    scale = total / float(sum(widths))
    twips = [max(1, int(w * scale)) for w in widths]

    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for column, w in zip(grid.findall(qn("w:gridCol")), twips):
            column.set(qn("w:w"), str(w))
    for tr in table._tbl.findall(qn("w:tr")):
        col = 0
        for tc in tr.findall(qn("w:tc")):
            tc_pr = tc.get_or_add_tcPr()
            grid_span = tc_pr.find(qn("w:gridSpan"))
            span = int(grid_span.get(qn("w:val"))) if grid_span is not None else 1
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.insert(0, tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(sum(twips[col:col + span])))
            col += span


def _build_picture(document, item: Mapping):
    try:
        data = base64.b64decode(item["data_b64"])
        section = document.sections[-1]
        text_width = int(section.page_width - section.left_margin - section.right_margin)
    except Exception:
        data, text_width = b"", int(Inches(6.5))
    if not data:
        return None
    natural = int(item.get("width") or 0) * 9525  # px at 96 dpi -> EMU
    width = Emu(min(natural, text_width) if natural else text_width)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(6)
    try:
        paragraph.add_run().add_picture(BytesIO(data), width=width)
    except Exception:
        paragraph._p.getparent().remove(paragraph._p)
        return None
    return paragraph._p
