"""Builders for fake score reports (test data only) used by the score-report tests."""
import io
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw

INDEX_ROWS = [
    ["Index", "Composite Score", "Percentile Rank", "95% Confidence Interval", "Qualitative Description"],
    ["Verbal Comprehension (VCI)", "104", "61", "97-110", "Average"],
    ["Visual Spatial (VSI)", "96", "39", "89-104", "Average"],
    ["Fluid Reasoning (FRI)", "88", "21", "82-96", "Low Average"],
    ["Working Memory (WMI)", "79", "8", "73-88", "Very Low"],
    ["Processing Speed (PSI)", "91", "27", "84-100", "Average"],
    ["Full Scale IQ (FSIQ)", "91", "27", "86-96", "Average"],
]
SUBTEST_ROWS = [
    ["Subtest", "Scaled Score", "Percentile Rank"],
    ["Similarities", "11", "63"],
    ["Vocabulary", "10", "50"],
    ["Block Design", "9", "37"],
    ["Digit Span", "6", "9"],
]


def bar_chart_png(width: int = 900, height: int = 480) -> bytes:
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((20, 10), "WISC-V Index Score Profile (test data)", fill="black")
    for i, value in enumerate([104, 96, 88, 79, 91]):
        x0 = 80 + i * 150
        draw.rectangle([x0, height - 40 - value * 3, x0 + 90, height - 40], fill=(60, 110, 190))
        draw.text((x0, height - 30), ["VCI", "VSI", "FRI", "WMI", "PSI"][i], fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def logo_png() -> bytes:
    image = Image.new("RGB", (60, 24), (200, 30, 30))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def make_score_report_docx(path: Path, *, with_graph: bool = True) -> Path:
    document = Document()
    document.add_picture(io.BytesIO(logo_png()), width=Inches(0.4))  # decorative logo: must be ignored
    document.add_paragraph("Score Report - Test Client (fake data)")
    document.add_paragraph("WISC-V Primary Index Scores")
    table = document.add_table(rows=len(INDEX_ROWS), cols=len(INDEX_ROWS[0]))
    for r, row in enumerate(INDEX_ROWS):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    document.add_paragraph("Subtest Scaled Scores")
    table = document.add_table(rows=len(SUBTEST_ROWS), cols=len(SUBTEST_ROWS[0]))
    for r, row in enumerate(SUBTEST_ROWS):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    if with_graph:
        document.add_paragraph("Index Score Profile")
        document.add_picture(io.BytesIO(bar_chart_png()), width=Inches(5.5))
        document.add_paragraph("Index Score Profile (repeated page header copy)")
        document.add_picture(io.BytesIO(bar_chart_png()), width=Inches(5.5))  # duplicate: must be de-duplicated
    document.save(str(path))
    return path


def _pdf_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_score_report_pdf(path: Path, *, rows=None) -> Path:
    """Hand-built one-page PDF with a ruled table, written without a PDF library
    so the tests need no extra dependencies."""
    rows = rows or INDEX_ROWS
    col_w, row_h, left, top = 100, 22, 40, 740
    cols = len(rows[0])
    ops = ["0.5 w"]
    for r in range(len(rows) + 1):
        y = top - r * row_h
        ops.append(f"{left} {y} m {left + cols * col_w} {y} l S")
    for c in range(cols + 1):
        x = left + c * col_w
        ops.append(f"{x} {top} m {x} {top - len(rows) * row_h} l S")
    ops.append("BT /F1 8 Tf")
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            ops.append(
                f"1 0 0 1 {left + c * col_w + 3} {top - r * row_h - 14} Tm ({_pdf_escape(value[:18])}) Tj"
            )
    ops.append("ET")
    content = "\n".join(ops).encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, xref)
    path.write_bytes(bytes(out))
    return path


def make_image_pdf(path: Path) -> Path:
    """One-page PDF holding just a raster graph (Pillow writes a valid image XObject)."""
    Image.open(io.BytesIO(bar_chart_png())).convert("RGB").save(str(path), format="PDF")
    return path
