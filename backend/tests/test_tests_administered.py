"""Tests that were not administered are taken out of the Word report (all names are made up)."""
import io
import json

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from services.docx_exporter import create_report_docx
from services.docx_test_results import _text
from services.test_pruning import families_in, prune_unadministered

CDBC = "SunnyHill CDBC Psychology Assessment - Boy"
PSYCHED = "Psycho-Educational Assessment - Boy"
ASD = "ASD Clinical Diagnostic Assessment Report"
METADATA = "# Report Metadata (Front Page)\nClient Full Name: Alex Test\n"


def given(test):
    return {"test_name": test, "items": []}


def export(title, *groups, content=METADATA, report=None, administered=None):
    stream = create_report_docx(
        title=title, patient_name="Alex Test", report_type="x", content=content,
        test_results=list(groups), fill_report=report, tests_administered=administered,
    )
    return Document(io.BytesIO(stream.getvalue()))


def paragraphs(doc):
    return [text for text in (_text(el) for el in doc.element.body if el.tag == qn("w:p")) if text]


def has(doc, startswith):
    return any(t.startswith(startswith) for t in paragraphs(doc))


# --------------------------------------------------------------------------- #
# recognising tests
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("first, second", [
    ("WIAT-4", "Wechsler Individual Achievement Test – 3rd Edition  (WIAT-III)"),
    ("ABAS-III", "Adaptive Behavior Assessment System- Third Edition (ABAS-3)"),
    ("Aseba", "Achenbach Scales"),
    ("Brief", "Behavior Rating Inventory of Executive Function, 2nd Edition (BRIEF-2)"),
    ("D-KEFS", "Delis–Kaplan Executive Function System (D–KEFS)-"),
    ("Scared", "Screen for Child Anxiety-Related Disorders (SCARED)"),
    ("Rey Complex Figure Test-1", "Rey–Osterrieth Complex Figure Test (RCFT)"),
    ("MASC", "Multidimensional Anxiety Scale for Children- Second Edition"),
])
def test_the_dropdown_name_and_the_templates_wording_are_the_same_test(first, second):
    assert families_in(first) & families_in(second)


def test_different_tests_are_not_confused():
    assert not families_in("WISC-V") & families_in("WAIS-IV")
    assert not families_in("BASC-3") & families_in("BRIEF-2")


# --------------------------------------------------------------------------- #
# the Tests Administered lists and the result blocks
# --------------------------------------------------------------------------- #
def test_only_the_tests_given_stay_in_the_tests_administered_list_and_results():
    doc = export(CDBC, given("WISC-V"), given("ABAS-III"))
    items = paragraphs(doc)

    assert "Wechsler Intelligence Scale for Children – 5th Edition (WISC-V)" in items
    assert "Adaptive Behavior Assessment System- Third Edition (ABAS-3)" in items
    for gone in (
        "Bayley Scales of Infant and Toddler Development-4 (Bayley-4)",
        "Wechsler Adult Intelligence Scale – 4th Edition (WAIS-IV)",
        "Behaviour Assessment System for Children- Third Edition (BASC-3)",
        "Vineland Adaptive Behavior Scales- Third Edition (VABS-III)",
    ):
        assert gone not in items
    # result blocks follow the same rule
    assert not has(doc, "Cognitive Score Summary Table: WAIS-IV")
    assert not has(doc, "WPPSI-IV Score Table")
    assert has(doc, "Cognitive Score Summary Table: WISC-5")
    assert has(doc, "Adaptive Functioning Score Summary Table: ABAS-3")


def test_section_headings_and_the_signatures_survive():
    doc = export(CDBC, given("WISC-V"))
    items = paragraphs(doc)
    for kept in ("TESTS ADMINISTERED", "General Cognitive Abilities", "Charmaine Miranda, Ph.D.", "(add your signature)"):
        assert kept in items
    assert any(t.startswith("Appendix 2") for t in items)


def test_headings_left_with_nothing_under_them_go_too():
    without = export(CDBC, given("WISC-V"))
    assert not has(without, "Attention, Activity and Executive Functioning")      # no tests under it were given
    assert not has(without, "PARENT/ TEACHER/SELF REPORT MEASURES ADMINISTERED")   # its whole list went

    with_brief = export(CDBC, given("WISC-V"), given("BRIEF-2"))
    assert has(with_brief, "Attention, Activity and Executive Functioning")
    assert has(with_brief, "PARENT/ TEACHER/SELF REPORT MEASURES ADMINISTERED")


def test_a_test_given_without_results_keeps_its_template_block_and_blank_tables():
    only_wisc = export(PSYCHED, given("WISC-V"))
    assert has(only_wisc, "WISC-V Index Scores Summary")
    assert not has(only_wisc, "WAIS-IV Index Scores Summary")
    assert len(only_wisc.tables) < len(export(PSYCHED).tables)   # the other tests' tables are gone


def test_nothing_is_removed_when_no_test_is_recorded():
    plain = export(PSYCHED)
    nameless = export(PSYCHED, {"test_name": "", "items": [{"kind": "table", "rows": [["a", "b"]]}]})
    assert len(plain.tables) == len(nameless.tables) and paragraphs(plain) == paragraphs(nameless)
    assert has(plain, "WAIS-IV Index Scores Summary")


def test_tests_listed_on_the_report_count_even_without_results_or_a_group():
    doc = export(PSYCHED, administered=["WISC-V"])
    assert has(doc, "WISC-V Index Scores Summary") and not has(doc, "WAIS-IV Index Scores Summary")


def test_a_test_table_section_in_the_report_counts_as_administered():
    content = METADATA + "# Test Table: WISC-V\nNotes about the WISC-V.\n"
    doc = export(PSYCHED, content=content)
    assert has(doc, "WISC-V Index Scores Summary") and not has(doc, "WAIS-IV Index Scores Summary")


def test_another_tests_name_on_a_table_caption_does_not_cost_a_given_test_its_table():
    """The Psycho-Ed template labels the WIAT table "Academic Scores - KTEA-3" under the WIAT text."""
    doc = export(PSYCHED, given("WIAT-IV"))
    assert has(doc, "Academic abilities were tested using the WIAT-III")
    assert has(doc, "Academic Scores")
    baseline = export(PSYCHED, given("WISC-V"))
    assert len(doc.tables) >= 1 and not has(baseline, "Academic abilities were tested using the WIAT-III")


def test_asd_template_lists_are_trimmed_too():
    doc = export(ASD, given("ABAS-III"))
    assert has(doc, "Adaptive Behavior Assessment System")
    assert not has(doc, "Autism Diagnostic Observation Schedule")


def test_what_was_removed_is_reported():
    report = []
    export(CDBC, given("WISC-V"), report=report)
    pruned = next(r for r in report if r["result"] == "pruned")
    assert any("Bayley" in label for label in pruned["removed"])
    assert not any("WISC-V" == label for label in pruned["removed"])
    assert report[-1] is pruned                      # per-test entries (when there are any) stay first


# --------------------------------------------------------------------------- #
# safety
# --------------------------------------------------------------------------- #
def _tree(*paragraphs_xml):
    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    return etree.fromstring(f"<w:document {ns}><w:body>{''.join(paragraphs_xml)}</w:body></w:document>")


def _p(text, section_break=False):
    brk = "<w:pPr><w:sectPr/></w:pPr>" if section_break else ""
    return f"<w:p>{brk}<w:r><w:t>{text}</w:t></w:r></w:p>"


def test_a_section_break_is_never_deleted_with_a_block():
    root = _tree(
        _p("Tests Administered"), _p("WISC-V"), _p("WAIS-IV"), _p("Test Results"),
        _p("WAIS-IV"), _p("The WAIS-IV is a measure of adult intelligence given to older teens and adults.", section_break=True),
        _p("Summary of Assessment Results"),
    )
    pruned = prune_unadministered(root, ["WISC-V"], ["Tests Administered", "Test Results"])
    texts = [_text(p) for p in root.iter(qn("w:p"))]
    assert pruned is not None
    assert any("The WAIS-IV is a measure" in t for t in texts)      # the paragraph holding the section break stays
    assert "WISC-V" in texts and "Summary of Assessment Results" in texts


def test_a_list_entry_we_cannot_place_is_kept():
    root = _tree(
        _p("Tests Administered"), _p("WISC-V"), _p("Some Brand New Questionnaire"), _p("WAIS-IV"),
        _p("Summary of Assessment Results"),
    )
    prune_unadministered(root, ["WISC-V"], ["Tests Administered"])
    texts = [_text(p) for p in root.iter(qn("w:p"))]
    assert "Some Brand New Questionnaire" in texts and "WAIS-IV" not in texts


def test_nothing_changes_if_no_given_test_can_be_recognised():
    root = _tree(_p("Tests Administered"), _p("WISC-V"), _p("WAIS-IV"), _p("Summary of Assessment Results"))
    assert prune_unadministered(root, ["Some Brand New Questionnaire"], []) is None
    assert [_text(p) for p in root.iter(qn("w:p"))].count("WAIS-IV") == 1


# --------------------------------------------------------------------------- #
# through the API: tests with no results are remembered
# --------------------------------------------------------------------------- #
@pytest.fixture()
def client_and_session():
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool
    from fastapi.testclient import TestClient

    from database.db import Base, get_db
    from main import app

    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(autocommit=False, autoflush=False, bind=engine)

    def override():
        db = Session()
        try:
            yield db
        finally:
            db.close()

    previous = app.dependency_overrides.get(get_db)
    app.dependency_overrides[get_db] = override
    try:
        yield TestClient(app), Session
    finally:
        if previous is None:
            app.dependency_overrides.pop(get_db, None)
        else:
            app.dependency_overrides[get_db] = previous


def test_api_saves_tests_administered_exports_without_the_rest_and_hides_empty_rows(client_and_session):
    client, Session = client_and_session
    from tests.test_score_report_api import _template_id

    template_id = _template_id(client)

    created = client.post("/api/reports/generate", json={
        "title": PSYCHED, "patient_name": "Alex Test", "report_type": "psychoeducational", "template_id": template_id,
        "additional_inputs": {"Report Metadata (Front Page)": "Client Full Name: Alex Test"},
        "tests_administered": ["WISC-V", "wisc-v", "  "],
        "test_results": [],
    })
    assert created.status_code == 200
    report_id = created.json()["id"]

    assert client.get(f"/api/reports/{report_id}/test-results").json() == []     # nothing attached: nothing listed

    exported = client.get(f"/api/reports/{report_id}/export-docx")
    assert exported.status_code == 200
    doc = Document(io.BytesIO(exported.content))
    assert has(doc, "WISC-V Index Scores Summary") and not has(doc, "WAIS-IV Index Scores Summary")
    summary = json.loads(exported.headers["x-score-fill"])
    assert any(entry["result"] == "pruned" for entry in summary)
