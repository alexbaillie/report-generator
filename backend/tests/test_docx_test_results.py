import base64
import io

from docx import Document
from docx.oxml.ns import qn

from services.docx_exporter import create_report_docx
from tests.score_report_samples import bar_chart_png

IMAGE_B64 = base64.b64encode(bar_chart_png()).decode()
METADATA = "# Report Metadata (Front Page)\nClient Full Name: Alex Test\n"
PSYCHED = "Psycho-Educational Assessment - Boy"
CDBC = "SunnyHill CDBC Psychology Assessment - Boy"
ASD = "ASD Clinical Diagnostic Assessment Report"
GENERIC = "Standard Intake Assessment"


def group(test, *, tables=True, image=True):
    items = []
    if tables:
        items.append({"kind": "table", "caption": f"[{test}] scores", "rows": [["Index", "Score"], [f"{test}-ROW", "100"]]})
    if image:
        items.append(
            {"kind": "image", "caption": f"[{test}] graph", "content_type": "image/png",
             "data_b64": IMAGE_B64, "width": 800, "height": 400}
        )
    return {"test_name": test, "items": items}


def export(title, *groups):
    stream = create_report_docx(
        title=title, patient_name="Alex Test", report_type="x", content=METADATA, test_results=list(groups)
    )
    return Document(io.BytesIO(stream.getvalue()))


def text(element):
    return "".join(t.text or "" for t in element.iter(qn("w:t"))).strip()


def paragraphs(doc):
    return [text(el) for el in doc.element.body if el.tag == qn("w:p") and text(el)]


def index_of(doc, startswith, after=0):
    items = paragraphs(doc)
    return next(i for i, t in enumerate(items) if i >= after and t.startswith(startswith))


def test_results_land_under_the_tests_own_heading_and_replace_its_blank_tables():
    baseline = export(PSYCHED)
    doc = export(PSYCHED, group("WISC-V"))

    items = paragraphs(doc)
    heading = index_of(doc, "WISC-V")
    next_test = index_of(doc, "WAIS-IV", after=heading)
    inserted = items.index("[WISC-V] scores")
    assert heading < inserted < next_test
    assert items.index("[WISC-V] graph") < next_test
    # the template's WISC-V description is still there, ahead of the new content
    assert any("Wechsler Intelligence Scale for Children (WISC-V)" in t for t in items[heading:inserted])
    # its seven blank score tables were replaced by the one uploaded table
    assert len(doc.tables) == len(baseline.tables) - 7 + 1
    assert len(doc.inline_shapes) == len(baseline.inline_shapes) + 1
    assert not any("insert WISC-V graph here" in t for t in items)


def test_a_neighbouring_tests_table_is_not_swallowed():
    """WRAML-3's block ends where CVLT-C's heading begins ("cvlt-c" must read as a test)."""
    doc = export(PSYCHED, group("WRAML-3"))

    assert any(t.startswith("CVLT-C Auditory Memory Scores") for t in paragraphs(doc))


def test_every_psyched_test_lands_in_its_own_block_in_template_order():
    tests = ["WISC-V", "WAIS-IV", "WRAML-3", "CVLT-C", "WIAT-IV", "BRIEF", "ABAS-III"]
    doc = export(PSYCHED, *[group(t, image=False) for t in tests])

    items = paragraphs(doc)
    positions = [items.index(f"[{t}] scores") for t in tests]
    assert positions == sorted(positions)  # same order as the template lays the tests out
    cvlt_heading = next(i for i, t in enumerate(items) if t.startswith("The California Verbal Learning Test"))
    assert cvlt_heading < items.index("[CVLT-C] scores") < index_of(doc, "Rey Complex Figure Test")


def test_roman_and_arabic_numerals_match_the_template_wording():
    # dropdown says "WIAT-IV" / "ABAS-III"; the template says "WIAT-III" / "ABAS-3"
    doc = export(PSYCHED, group("WIAT-IV", image=False), group("ABAS-III", image=False))

    items = paragraphs(doc)
    academic = next(i for i, t in enumerate(items) if t.startswith("Academic abilities were tested using the WIAT-III"))
    assert items.index("[WIAT-IV] scores") > academic
    adaptive = next(
        i for i, t in enumerate(items) if t.startswith("Alex") and "Adaptive Behavior Assessment" in t
    )
    assert items.index("[ABAS-III] scores") > adaptive


def test_tables_only_keeps_the_graph_placeholder_and_graph_only_keeps_blank_tables():
    tables_only = export(PSYCHED, group("WISC-V", image=False))
    graph_only = export(PSYCHED, group("WISC-V", tables=False))
    baseline = export(PSYCHED)

    assert any("insert WISC-V graph here" in t for t in paragraphs(tables_only))
    assert not any("insert WISC-V graph here" in t for t in paragraphs(graph_only))
    assert len(graph_only.tables) == len(baseline.tables)  # blank shells untouched


def test_a_test_with_no_block_in_the_template_gets_its_own_block_before_the_summary():
    doc = export(PSYCHED, group("Vineland-3"))

    items = paragraphs(doc)
    heading = items.index("Vineland-3")
    assert heading < items.index("[Vineland-3] scores") < index_of(doc, "Summary of Assessment Results")


def test_cdbc_and_asd_templates_are_supported():
    cdbc = export(CDBC, group("WISC-V"), group("ABAS-III"))
    asd = export(ASD, group("ABAS-III"), group("WISC-V"))

    cdbc_items, asd_items = paragraphs(cdbc), paragraphs(asd)
    assert cdbc_items.index("[WISC-V] scores") > index_of(cdbc, "WISC-V")
    assert asd_items.index("[ABAS-III] scores") > index_of(asd, "Adaptive Behaviour")
    # not in the ASD template: its own block, ahead of the closing recommendations
    closing = index_of(asd, "Additional Recommendations", after=index_of(asd, "Tests Administered"))
    assert asd_items.index("[WISC-V] scores") < closing


def test_the_same_test_attached_twice_is_merged_into_one_block():
    first = group("WISC-V", image=False)
    second = {"test_name": "wisc-v", "items": [{"kind": "table", "caption": "[second] table", "rows": [["a", "b"], ["x", "y"]]}]}

    items = paragraphs(export(PSYCHED, first, second))

    assert items.index("[WISC-V] scores") < items.index("[second] table")
    assert items.count("WISC-V") == 1  # no duplicate fallback heading


def test_generic_export_appends_a_labelled_block():
    doc = export(GENERIC, group("WISC-V"))

    items = paragraphs(doc)
    assert items.index("WISC-V") < items.index("[WISC-V] scores")
    assert len(doc.tables) == 1 and len(doc.inline_shapes) == 1


def test_no_results_leaves_the_export_untouched():
    plain = create_report_docx(title=PSYCHED, patient_name="Alex Test", report_type="x", content=METADATA)
    with_empty = create_report_docx(
        title=PSYCHED, patient_name="Alex Test", report_type="x", content=METADATA,
        test_results=[{"test_name": "WISC-V", "items": []}, {"test_name": "", "items": [{"kind": "table", "rows": [["a", "b"]]}]}],
    )

    assert plain.getvalue() == with_empty.getvalue()


def test_inserted_tables_keep_word_schema_order_and_are_separated():
    doc = export(PSYCHED, group("WISC-V", image=False), group("WAIS-IV", image=False))

    order = ["tblStyle", "tblW", "jc", "tblBorders", "tblLayout", "tblCellMar", "tblLook"]
    for table in doc.tables:
        names = [child.tag.split("}")[1] for child in table._tbl.tblPr]
        known = [n for n in names if n in order]
        assert known == sorted(known, key=order.index), names
    body = list(doc.element.body)
    for left, right in zip(body, body[1:]):
        assert not (left.tag == qn("w:tbl") and right.tag == qn("w:tbl"))  # Word would merge them
