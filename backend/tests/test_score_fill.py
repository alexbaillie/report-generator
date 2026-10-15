"""Filling the Word template's own score tables from a score report (all numbers are fake)."""
import base64
import io
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from services.docx_exporter import create_report_docx
from services.docx_test_results import _build_table
from services.score_fill import parse_source_tables, split_label
from services.score_report_extractor import extract_score_report
from tests.score_report_samples import bar_chart_png

METADATA = "# Report Metadata (Front Page)\nClient Full Name: Alex Test\n"
PSYCHED = "Psycho-Educational Assessment - Boy"
CDBC = "SunnyHill CDBC Psychology Assessment - Boy"

INDEX_HEAD = ["Composite", "Code", "Sum of Scaled Scores", "Composite Score", "Percentile Rank",
              "95% Confidence Interval", "Qualitative Description", "SEM"]
SUBTEST_HEAD = ["Subtest", "Code", "Total Raw Score", "Scaled Score", "Percentile Rank", "Age Equivalent", "SEM"]


def pearson_index_table():
    cols = len(INDEX_HEAD)
    return {
        "kind": "table", "caption": "Primary Index Scores",
        "rows": [
            ["Primary Index Scores"] + [""] * (cols - 1),
            INDEX_HEAD,
            ["Verbal Comprehension", "VCI", "18", "101", "53", "94-108", "Average", "4.7"],
            ["Visual Spatial", "VSI", "19", "104", "61", "96-111", "Average", "4.5"],
            ["Fluid Reasoning", "FRI", "20", "99", "47", "92-106", "Average", "3.3"],
            ["Working Memory", "WMI", "-", "-", "-", "-", "-", "-"],
            ["Processing Speed", "PSI", "17", "88", "21", "81-97", "Low Average", "4.2"],
            ["Full Scale IQ", "FSIQ", "48", "97", "42", "92-102", "Average", "3.0"],
        ],
        "merges": [[0, 0, 1, cols]],
        "widths": [1700, 600, 900, 900, 900, 1100, 1500, 600],
    }


def pearson_subtest_table():
    cols = len(SUBTEST_HEAD)
    return {
        "kind": "table", "caption": "Primary Subtest Scores",
        "rows": [
            ["Primary Subtest Scores"] + [""] * (cols - 1),
            SUBTEST_HEAD,
            ["Similarities", "SI", "30", "11", "63", "14:2", "1.3"],
            ["Vocabulary", "VC", "33", "10", "50", "13:6", "1.1"],
            ["Digit Span", "DS", "23", "-", "-", "-", "-"],
            ["Coding", "CD", "39", "7", "16", "9:10", "1.1"],
        ],
        "merges": [[0, 0, 1, cols]],
    }


def cover_png():
    from PIL import Image
    buffer = io.BytesIO()
    Image.new("RGB", (400, 200), (30, 60, 160)).save(buffer, format="PNG")
    return buffer.getvalue()


def graph():
    return {"kind": "image", "caption": "Index graph", "content_type": "image/png",
            "data_b64": base64.b64encode(bar_chart_png()).decode(), "width": 800, "height": 400}


def export(title, *groups, report=None):
    stream = create_report_docx(
        title=title, patient_name="Alex Test", report_type="x", content=METADATA,
        test_results=list(groups), fill_report=report,
    )
    return Document(io.BytesIO(stream.getvalue()))


def rows_with(doc, label, exact=False):
    """Every table row whose first cell is (or starts with) ``label``. The templates
    carry several tests' tables, so the same label can appear more than once."""
    found = []
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0].text.strip()
            if first == label if exact else first.startswith(label):
                found.append([c.text.strip() for c in row.cells])
    return found


def filled_row(doc, label, exact=False):
    """The one row for ``label`` that has scores in it (the score columns come last);
    every other copy in the template must still be blank."""
    rows = rows_with(doc, label, exact)
    filled = [r for r in rows if r[-1]]
    assert len(filled) == 1, f"{label!r}: {rows}"
    return filled[0]


def table_count(doc):
    return len(doc.tables)


# --------------------------------------------------------------------------- #
# reading the score report
# --------------------------------------------------------------------------- #
def test_labels_split_into_name_and_code():
    assert split_label("Similarities (SI)") == ("similarities", "SI")
    assert split_label("Full Scale (FSIQ)*") == ("full scale", "FSIQ")
    assert split_label("Letter-Number Sequencing (LNS)***") == ("letter number sequencing", "LN")
    assert split_label("Cancellation") == ("cancellation", "")


def test_source_records_skip_title_rows_and_flag_unadministered():
    records = {r.code: r for r in parse_source_tables([pearson_index_table(), pearson_subtest_table()])}
    assert records["VCI"].fields["standard"] == "101"
    assert records["VCI"].fields["ci"] == "94-108"
    assert records["SI"].fields["scaled"] == "11"
    assert records["WMI"].administered is False
    assert records["DS"].administered is False
    assert "Primary Index Scores" not in {r.display for r in records.values()}


# --------------------------------------------------------------------------- #
# filling the template
# --------------------------------------------------------------------------- #
def test_fill_keeps_the_templates_tables_and_fills_only_the_blank_score_cells():
    baseline = export(PSYCHED)
    report = []
    doc = export(PSYCHED, {"test_name": "WISC-V", "items": [pearson_index_table(), pearson_subtest_table()]}, report=report)

    assert table_count(doc) == table_count(baseline)  # nothing added, nothing replaced
    # template wording is untouched, scores are in the template's own columns
    assert filled_row(doc, "Verbal Comprehension (VCI)")[:5] == ["Verbal Comprehension (VCI)", "101", "53", "94-108", "Average"]
    assert filled_row(doc, "Full Scale (FSIQ)")[1:] == ["97", "42", "92-102", "Average"]
    similarities = filled_row(doc, "Similarities (SI)")
    assert similarities[-2:] == ["11", "63"]   # scaled score, percentile rank; skills text left alone
    assert similarities[1].startswith("Verbal abstract reasoning")

    assert report[0]["test_name"] == "WISC-V"
    assert report[0]["result"] == "filled" and report[0]["filled"] >= 12


def test_scores_marked_with_a_dash_say_not_assessed_and_are_reported():
    report = []
    doc = export(PSYCHED, {"test_name": "WISC-V", "items": [pearson_index_table(), pearson_subtest_table()]}, report=report)
    # Working Memory was not administered: its template row says so, with no numbers
    marked = [r for r in rows_with(doc, "Working Memory (WMI)") if "Not assessed" in r]
    assert len(marked) == 1 and marked[0][1:] == ["Not assessed", "", "", ""]
    assert any("Working Memory" in n for n in report[0]["not_administered"])
    assert report[0]["marked_not_assessed"] >= 1


def test_only_blank_cells_are_filled_never_existing_text():
    baseline = export(PSYCHED)
    table = next(t for t in baseline.tables if t.rows[1].cells[0].text.startswith("Verbal Comprehension (VCI)"))
    assert [c.text for c in table.rows[1].cells[1:]] == ["", "", "", ""]  # blank in the template

    # a template row that already says something keeps it
    from services.score_fill import fill_template_tables
    cell = table.rows[1].cells[1]
    cell.paragraphs[0].add_run("keep me")
    stats = fill_template_tables([table._tbl], parse_source_tables([pearson_index_table()]))
    assert cell.text == "keep me"
    assert stats.filled >= 3  # the other cells in the row still fill


def test_conflicting_duplicates_are_left_blank():
    clash = pearson_index_table()
    other = pearson_index_table()
    other["rows"][2][3] = "55"  # same code, different composite score
    report = []
    doc = export(PSYCHED, {"test_name": "WISC-V", "items": [clash, other, pearson_subtest_table()]}, report=report)
    assert all(r[1:] == ["", "", "", ""] for r in rows_with(doc, "Verbal Comprehension (VCI)"))
    assert any("Verbal Comprehension" in a for a in report[0]["ambiguous"])


def test_a_column_is_only_filled_from_the_same_kind_of_score():
    # The report gives a *standard* score where the template column wants a *scaled* score
    wrong_kind = {
        "kind": "table", "caption": "",
        "rows": [["Subtest", "Code", "Standard Score", "Percentile Rank"],
                 ["Similarities", "SI", "101", "53"]],
    }
    doc = export(PSYCHED, {"test_name": "WISC-V", "items": [wrong_kind]})
    assert not any("101" in r for r in rows_with(doc, "Similarities (SI)"))   # never into a scaled-score column
    filled = [r for r in rows_with(doc, "Similarities (SI)") if r[-1]]
    assert len(filled) == 1 and filled[0][-1] == "53" and filled[0][-2] == ""  # the percentile still fills


def test_graph_replaces_the_placeholder_while_tables_are_filled():
    baseline = export(PSYCHED)
    doc = export(PSYCHED, {"test_name": "WISC-V", "items": [pearson_index_table(), graph()]})
    assert len(doc.inline_shapes) == len(baseline.inline_shapes) + 1
    assert table_count(doc) == table_count(baseline)
    assert filled_row(doc, "Verbal Comprehension (VCI)")[1] == "101"


def test_nothing_matched_falls_back_to_inserting_the_uploaded_tables():
    stranger = {"kind": "table", "caption": "Odd scores",
                "rows": [["Thing", "Percentile Rank", "Scaled Score"], ["Widgets", "50", "10"]]}
    report = []
    doc = export(PSYCHED, {"test_name": "WISC-V", "items": [stranger]}, report=report)
    assert rows_with(doc, "Widgets") == [["Widgets", "50", "10"]]
    assert report[0]["result"] == "replaced"


def test_insert_mode_replaces_the_templates_tables_with_the_report_as_is():
    baseline = export(PSYCHED)
    report = []
    doc = export(PSYCHED, {"test_name": "WISC-V", "mode": "insert",
                           "items": [pearson_index_table()]}, report=report)
    assert report[0]["result"] == "replaced"
    assert table_count(doc) < table_count(baseline) + 1
    assert rows_with(doc, "Verbal Comprehension", exact=True)[0][:2] == ["Verbal Comprehension", "VCI"]
    # the title row is merged across the table and isn't shaded as the header
    title_table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Primary Index Scores")
    assert len({id(c._tc) for c in title_table.rows[0].cells}) == 1


def test_other_templates_are_filled_the_same_way():
    report = []
    doc = export(CDBC, {"test_name": "WISC-V", "items": [pearson_index_table(), pearson_subtest_table()]}, report=report)
    assert report[0]["result"] in ("filled", "replaced")
    if report[0]["result"] == "filled":
        assert any("101" in r for r in rows_with(doc, "Verbal Comprehension"))


def test_a_test_with_no_template_block_is_still_added():
    report = []
    doc = export(PSYCHED, {"test_name": "ZZTEST-9", "items": [pearson_index_table()]}, report=report)
    assert report[0]["result"] == "added"
    assert rows_with(doc, "Verbal Comprehension", exact=True)[0][:2] == ["Verbal Comprehension", "VCI"]


# --------------------------------------------------------------------------- #
# fidelity of an inserted table
# --------------------------------------------------------------------------- #
def build(rows, merges=None, widths=None):
    doc = Document()
    tbl = _build_table(doc, rows, None, merges, widths)
    doc.element.body.insert(0, tbl)
    return Document(io.BytesIO(_save(doc))).tables[0]


def _save(doc):
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_inserted_table_merges_cells_and_keeps_widths():
    table = build(pearson_index_table()["rows"], pearson_index_table()["merges"], [1, 1, 1, 1, 1, 1, 1, 3])
    assert len({id(c._tc) for c in table.rows[0].cells}) == 1    # title row is one merged cell
    widths = [int(g.get(qn("w:w"))) for g in table._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))]
    assert widths[-1] > 2 * widths[0]


def test_title_row_is_not_shaded_as_the_header():
    table = build(pearson_index_table()["rows"], pearson_index_table()["merges"])

    def fill(row):
        shd = table.rows[row].cells[0]._tc.tcPr.find(qn("w:shd")) if table.rows[row].cells[0]._tc.tcPr is not None else None
        return None if shd is None else shd.get(qn("w:fill"))

    assert fill(0) is None
    assert fill(1) == "D9E2F3"


def test_plain_table_still_shades_its_first_row():
    table = build([["Index", "Score"], ["A", "1"]])
    shd = table.rows[0].cells[0]._tc.tcPr.find(qn("w:shd"))
    assert shd.get(qn("w:fill")) == "D9E2F3"


# --------------------------------------------------------------------------- #
# the extractor's side of the fidelity work
# --------------------------------------------------------------------------- #
def test_extractor_reports_merges_widths_and_flags_cover_pictures(tmp_path: Path):
    doc = Document()
    doc.add_picture(io.BytesIO(cover_png()))         # cover art, before any table
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Title"
    for c, text in enumerate(["Name", "Scaled Score", "Percentile Rank"]):
        table.cell(1, c).text = text
    for c, text in enumerate(["Similarities (SI)", "9", "37"]):
        table.cell(2, c).text = text
    doc.add_picture(io.BytesIO(bar_chart_png()))      # a real graph, after the table
    path = tmp_path / "report.docx"
    doc.save(path)

    items = extract_score_report(path)["items"]
    pictures = [i for i in items if i["kind"] == "image"]
    tables = [i for i in items if i["kind"] == "table"]
    assert pictures[0].get("suggested") is False
    assert pictures[1].get("suggested", True) is True
    assert tables[0]["merges"] == [[0, 0, 1, 3]]
    assert len(tables[0]["widths"]) == 3


# --------------------------------------------------------------------------- #
# a fully generated report: the template's headings are gone, its tables remain
# --------------------------------------------------------------------------- #
GENERATED = (
    METADATA
    + "# Cognitive & Neuropsychological Results\nThe WISC-V results show average ability.\n"
    + "# Academic Achievement\nReading is at grade level.\n"
    + "# Adaptive, Behavioural & Social-Emotional Functioning\nAdaptive skills are typical.\n"
)


def export_generated(title, *groups, report=None):
    stream = create_report_docx(
        title=title, patient_name="Alex Test", report_type="x", content=GENERATED,
        test_results=list(groups), fill_report=report,
    )
    return Document(io.BytesIO(stream.getvalue()))


def test_only_the_named_versions_tables_are_filled_when_the_templates_headings_are_gone():
    from services.docx_test_results import _norm, _text
    baseline = export_generated(PSYCHED)
    # sanity: the headings really are gone in this scenario
    assert not any(_norm(_text(el)) == "wisc5" for el in baseline.element.body)

    report = []
    doc = export_generated(PSYCHED, {"test_name": "WISC-V", "items": [pearson_index_table(), pearson_subtest_table()]}, report=report)
    assert report[0]["result"] == "filled" and report[0]["filled"] >= 12

    vci_rows = [[c.text.strip() for c in r.cells] for t in doc.tables for r in t.rows
                if r.cells[0].text.strip().startswith("Verbal Comprehension (VCI)")]
    assert sum(1 for r in vci_rows if r[1] == "101") == 1            # WISC-V only
    assert sum(1 for r in vci_rows if not any(r[1:])) == len(vci_rows) - 1   # WISC-IV / WAIS-IV stay blank
    assert table_count(doc) == table_count(baseline)


def test_a_request_for_another_version_never_writes_into_tables_headed_with_a_different_version():
    doc = export_generated(PSYCHED, {"test_name": "WISC-IV", "items": [pearson_index_table()]})
    for table in doc.tables:
        header = table.rows[0].cells[0].text.replace(chr(10), " ")
        if header.startswith(("WISC-V", "WAIS-IV")):
            assert not any(c.text.strip() == "101" for row in table.rows for c in row.cells), header
